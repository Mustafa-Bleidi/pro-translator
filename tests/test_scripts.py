"""Tests for the helper scripts.  Run:  python3 -m unittest discover -s tests -v"""
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SCRIPTS = os.path.join(ROOT, "pro-translator", "scripts")
FIX = os.path.join(ROOT, "tests", "fixtures")


def run(script, *args):
    p = subprocess.run([sys.executable, os.path.join(SCRIPTS, script), *args],
                       capture_output=True, text=True)
    return p.returncode, p.stdout + p.stderr


def load(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def dump(obj, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False)


def have(mod):
    try:
        __import__(mod)
        return True
    except ImportError:
        return False


class QACheck(unittest.TestCase):
    def test_good_translation_passes(self):
        code, out = run("qa_check.py", f"{FIX}/source.en.md", f"{FIX}/target.ar.md", "--target", "ar")
        self.assertEqual(code, 0, out)

    def test_broken_translation_is_caught(self):
        code, out = run("qa_check.py", f"{FIX}/source.en.md", f"{FIX}/target.ar.broken.md", "--target", "ar")
        self.assertEqual(code, 2, out)
        for needle in ("2026", "URL", "inline code", "code block", "placeholder", "untranslated", "list items"):
            self.assertIn(needle, out)


class Glossary(unittest.TestCase):
    def test_glossary(self):
        code, _ = run("glossary_check.py", f"{FIX}/glossary.csv", f"{FIX}/source.en.md", f"{FIX}/target.ar.md", "--target", "ar")
        self.assertEqual(code, 0)
        code, out = run("glossary_check.py", f"{FIX}/glossary.csv", f"{FIX}/source.en.md", f"{FIX}/target.ar.broken.md", "--target", "ar")
        self.assertEqual(code, 1, out)


class Preflight(unittest.TestCase):
    def test_runs(self):
        code, out = run("preflight.py", "--target", "ar", "--json")
        self.assertEqual(code, 0)
        self.assertIn("modules", json.loads(out))


@unittest.skipUnless(have("lxml") and have("docx") and have("pptx") and have("openpyxl"), "needs lxml, python-docx, python-pptx, openpyxl")
class OfficeRoundTrip(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def path(self, name):
        return os.path.join(self.tmp, name)

    def roundtrip(self, src, translate, out, *extra):
        seg = self.path("seg.json")
        code, msg = run("office_segments.py", "extract", src, "-o", seg)
        self.assertEqual(code, 0, msg)
        segs = load(seg)["segments"]
        tr = {s["id"]: translate(s["text"]) for s in segs}
        trp = self.path("tr.json")
        dump(tr, trp)
        code, msg = run("office_segments.py", "apply", src, trp, "-o", out, *extra)
        self.assertEqual(code, 0, msg)
        return segs

    def test_docx(self):
        import docx
        from docx.oxml import OxmlElement
        d = docx.Document()
        d.add_heading("Annual report", 1)
        p = d.add_paragraph("Revenue grew ")
        r = p.add_run("12.5%")
        r.bold = True
        p.add_run(" this year.")
        p2 = d.add_paragraph("Name:")
        p2.add_run().add_tab()
        p2.add_run("Value")
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Metric", "Total"
        d.sections[0].header.paragraphs[0].text = "Internal"
        src = self.path("a.docx")
        d.save(src)
        out = self.path("a_AR.docx")
        segs = self.roundtrip(src, lambda s: s.replace("Revenue grew ", "نمت الإيرادات ")
                              .replace(" this year.", " هذا العام.").replace("Annual report", "التقرير السنوي")
                              .replace("Name:", "الاسم:").replace("Value", "القيمة").replace("Metric", "المقياس")
                              .replace("Total", "الإجمالي").replace("Internal", "داخلي"), out, "--lang", "ar")
        self.assertTrue(any("<x1/>" in s["text"] for s in segs))
        o = docx.Document(out)
        texts = [p.text for p in o.paragraphs]
        self.assertIn("نمت الإيرادات 12.5% هذا العام.", texts)
        self.assertIn("الاسم:\tالقيمة", texts)
        self.assertTrue(o.paragraphs[1].runs[1].bold)
        self.assertEqual(o.tables[0].cell(0, 1).text, "الإجمالي")
        self.assertEqual(o.sections[0].header.paragraphs[0].text, "داخلي")
        xml = zipfile.ZipFile(out).read("word/document.xml").decode()
        self.assertIn("<w:bidi/>", xml)
        self.assertIn("<w:bidiVisual/>", xml)
        self.assertIn('w:bidi="ar-SA"', xml)
        # input untouched
        self.assertIn("Revenue grew", docx.Document(src).paragraphs[1].text)

    def test_missing_object_tag_is_reported_and_restored(self):
        import docx
        d = docx.Document()
        p = d.add_paragraph("A")
        p.add_run().add_tab()
        p.add_run("B")
        src = self.path("b.docx")
        d.save(src)
        seg = self.path("s.json")
        run("office_segments.py", "extract", src, "-o", seg)
        sid = load(seg)["segments"][0]["id"]
        trp = self.path("t.json")
        dump({sid: "<g1>X</g1><g2>Y</g2>"}, trp)
        code, out = run("office_segments.py", "apply", src, trp, "-o", self.path("b2.docx"))
        self.assertEqual(code, 2)
        self.assertIn("CRITICAL", out)
        self.assertIn("\t", docx.Document(self.path("b2.docx")).paragraphs[0].text)

    def test_pptx(self):
        import pptx
        prs = pptx.Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = "Welcome"
        s.placeholders[1].text_frame.text = "Fast setup"
        src = self.path("a.pptx")
        prs.save(src)
        out = self.path("a_DE.pptx")
        self.roundtrip(src, lambda s: {"Welcome": "Willkommen", "Fast setup": "Schnelle Einrichtung"}[s], out, "--lang", "de")
        texts = [sh.text_frame.text for sh in pptx.Presentation(out).slides[0].shapes if sh.has_text_frame]
        self.assertEqual(texts, ["Willkommen", "Schnelle Einrichtung"])

    def test_xlsx(self):
        import openpyxl
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"], ws["B1"], ws["C1"] = "Product", 12.5, "=B1*2"
        src = self.path("a.xlsx")
        wb.save(src)
        out = self.path("a_AR.xlsx")
        self.roundtrip(src, lambda s: {"Product": "المنتج"}[s], out, "--lang", "ar", "--rtl-sheets")
        ws2 = openpyxl.load_workbook(out).active
        self.assertEqual((ws2["A1"].value, ws2["B1"].value, ws2["C1"].value), ("المنتج", 12.5, "=B1*2"))
        self.assertTrue(ws2.sheet_view.rightToLeft)


if __name__ == "__main__":
    unittest.main()
